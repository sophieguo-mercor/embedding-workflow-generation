"""
Convert HDBSCAN config JSON artifacts into Excel workbooks (one .xlsx per config).

Each artifact written by the sweep (results/hdbscan/configs/<name>.json) becomes
results/hdbscan/configs_xlsx/<name>.xlsx with these sheets:

  workflows  — the §2.8 MECE consolidation: one row per merged workflow (category →
               workflow with volume rollups) + a residual Unclassified row. Only
               present when the artifact carries a `consolidation` block.
  clusters   — one row per raw micro-cluster (the naming + coverage table)
  summary    — config weights, min_cluster_size, n_clusters, coverage, noise mass,
               UMAP/backend settings, the LLM judge scores, and the consolidation
               summary metrics

The bulky per-ticket `labels` array is intentionally dropped.

Usage:
    python to_excel_hdbscan.py                                   # convert every config
    python to_excel_hdbscan.py --config raw_categorical__mcs30   # convert just one
    python to_excel_hdbscan.py --in-dir results/hdbscan/configs --out-dir results/hdbscan/configs_xlsx
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

import config as C

# Classic Excel conditional-format colors: fill + matching font.
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREEN_FONT = Font(color="006100")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006")

# Column order for the clusters sheet — readable-first, verbose fields last.
CLUSTER_COLUMNS = [
    "cluster_id", "title", "category", "qualifies", "coherent", "big_enough",
    "n_tickets", "hours", "hours_frac", "coherence_note", "keywords", "examples",
    "description",
]


def _clusters_frame(clusters: list[dict]) -> pd.DataFrame:
    """Cluster rows as a DataFrame, with `examples`/`keywords` flattened to text."""
    rows = []
    for c in clusters:
        row = dict(c)
        for list_col in ("examples", "keywords"):
            val = row.get(list_col, [])
            if isinstance(val, list):
                row[list_col] = "\n".join(str(e) for e in val)
        rows.append(row)
    df = pd.DataFrame(rows)
    # Show known columns first (those present), then any extras the pipeline adds.
    ordered = [col for col in CLUSTER_COLUMNS if col in df.columns]
    extras = [col for col in df.columns if col not in ordered]
    return df[ordered + extras]


def _workflows_frame(consolidation: dict) -> pd.DataFrame:
    """The §2.8 MECE consolidation as one row per workflow (category → workflow with
    volume rollups), plus a trailing Unclassified row. member_cluster_ids flattened
    to text so the row traces back to the clusters sheet."""
    rows = []
    cats = sorted(consolidation.get("taxonomy", []),
                  key=lambda c: -(c.get("rollup", {}).get("hours", 0)))
    for cat in cats:
        wfs = sorted(cat.get("workflows", []),
                     key=lambda w: -(w.get("rollup", {}).get("hours", 0)))
        for wf in wfs:
            r = wf.get("rollup", {})
            rows.append({
                "category": cat.get("category", ""),
                "workflow": wf.get("name", ""),
                "n_clusters": r.get("n_clusters", 0),
                "n_tickets": r.get("n_tickets", 0),
                "hours": r.get("hours", 0.0),
                "description": wf.get("description", ""),
                "merge_note": wf.get("merge_note", ""),
                "member_cluster_ids": ", ".join(str(i) for i in wf.get("member_cluster_ids", [])),
            })
    unc = consolidation.get("unclassified", {})
    ur = unc.get("rollup", {})
    rows.append({
        "category": "Unclassified",
        "workflow": "(residual)",
        "n_clusters": ur.get("n_clusters", 0),
        "n_tickets": ur.get("n_tickets", 0),
        "hours": ur.get("hours", 0.0),
        "description": unc.get("note", ""),
        "merge_note": "",
        "member_cluster_ids": ", ".join(str(i) for i in unc.get("member_cluster_ids", [])),
    })
    return pd.DataFrame(rows)


def _summary_frame(artifact: dict) -> pd.DataFrame:
    """Flattened config-level metadata as key/value rows."""
    items: list[tuple[str, object]] = [
        ("config", artifact.get("config")),
        ("min_cluster_size", artifact.get("min_cluster_size")),
        ("n_clusters", artifact.get("n_clusters")),
        ("noise_mass", artifact.get("noise_mass")),
    ]
    for key, val in (artifact.get("weights") or {}).items():
        items.append((f"weight.{key}", val))
    for key, val in (artifact.get("coverage") or {}).items():
        items.append((f"coverage.{key}", val))
    for key, val in (artifact.get("llm_score") or {}).items():
        items.append((f"llm_score.{key}", val))
    for key, val in (artifact.get("backends") or {}).items():
        items.append((f"backend.{key}", val))
    for key, val in (artifact.get("umap") or {}).items():
        items.append((f"umap.{key}", val))
    csol = artifact.get("consolidation") or {}
    if csol:
        items.append(("consolidation.model", csol.get("model")))
        for key, val in (csol.get("summary") or {}).items():
            items.append((f"consolidation.{key}", val))
    return pd.DataFrame(items, columns=["metric", "value"])


def _paint(cell, ok: bool) -> None:
    """Green cell for a passing value, red for a failing one."""
    cell.fill = GREEN_FILL if ok else RED_FILL
    cell.font = GREEN_FONT if ok else RED_FONT


def _color_clusters_sheet(ws) -> None:
    """Red/green the boolean cells and the hours_frac mass gate in-place."""
    headers = {cell.value: cell.column for cell in ws[1]}
    frac_col = headers.get("hours_frac")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, bool):  # qualifies / coherent / big_enough
                _paint(cell, cell.value)
            elif cell.column == frac_col and isinstance(cell.value, (int, float)):
                _paint(cell, cell.value >= C.MIN_CLUSTER_MASS_FRAC)


def config_to_excel(json_path: Path, out_dir: Path) -> Path:
    """Convert one HDBSCAN config artifact to an .xlsx; returns the written path."""
    artifact = json.loads(json_path.read_text(encoding="utf-8"))
    out_path = out_dir / f"{json_path.stem}.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # workflows first — the §2.8 consolidation is the headline view when present.
        consolidation = artifact.get("consolidation")
        if consolidation:
            _workflows_frame(consolidation).to_excel(
                writer, sheet_name="workflows", index=False)
        _clusters_frame(artifact.get("clusters", [])).to_excel(
            writer, sheet_name="clusters", index=False)
        _summary_frame(artifact).to_excel(
            writer, sheet_name="summary", index=False)
        _color_clusters_sheet(writer.sheets["clusters"])
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser(description="Convert HDBSCAN config JSON artifacts to Excel.")
    ap.add_argument("--config", default=None,
                    help="convert only this config (matches results/hdbscan/configs/<name>.json)")
    ap.add_argument("--in-dir", default=f"{C.RESULTS_DIR}/hdbscan/configs",
                    help="directory of HDBSCAN config JSON artifacts")
    ap.add_argument("--out-dir", default=f"{C.RESULTS_DIR}/hdbscan/configs_xlsx",
                    help="directory to write .xlsx files into")
    args = ap.parse_args()

    in_dir = Path(args.in_dir)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.config:
        paths = [in_dir / f"{args.config}.json"]
        if not paths[0].exists():
            raise SystemExit(f"No config artifact at {paths[0]}. Run the sweep first.")
    else:
        paths = sorted(in_dir.glob("*.json"))
        if not paths:
            raise SystemExit(f"No config artifacts found in {in_dir}. Run the sweep first.")

    for p in paths:
        out = config_to_excel(p, out_dir)
        print(f"  {p.name}  →  {out}")
    print(f"Wrote {len(paths)} workbook(s) to {out_dir}")


if __name__ == "__main__":
    main()
