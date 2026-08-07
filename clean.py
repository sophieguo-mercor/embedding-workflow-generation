#!/usr/bin/env python3
"""
Section 1.1 — Text cleaning for ticket descriptions and time-entry notes.

The raw `description` column is an HTML-to-text email dump: double-encoded XML
escapes (`_x000D_`), inline-image refs (`[cid:...]`), tracking-link wrappers
(`text<https://…exclaimer.net…>`), signature blocks, quoted reply threads, and
an "Incoming Email Processor" footer. Measured on a 20k-row sample, that noise
is 60–75% of the bytes; the actual signal — the requester's ask / symptom /
affected system — is a short block at the top (median ~274 chars). We strip to
that block.

Time-entry `summarynotes` / `internalnotes` are already clean engineer prose
(the same noise patterns are <0.5% there), so they get the same structural
scrub but KEEP urls — a `learn.microsoft.com` link in an internal note is
signal, not boilerplate.

No PII redaction happens here (by design): the cleaned twins keep full fidelity
for reuse, and `pii.py` still gates PII at publish time (§5.1).

Everything is pure regex/string transformation — no eval, no code execution,
no network. `clean_workbook` streams the large (~66–82MB) exports with openpyxl
in read-only/write-only mode and writes cleaned twins with identical schema.

CLI:
    python clean.py                       # clean both exports → data/cleaned/
    python clean.py --tickets X --time-entries Y --out-dir data/cleaned
"""
from __future__ import annotations

import re

import config as C

# ── individual scrubbers ──────────────────────────────────────────────────────

# Double-encoded XML escapes leaked by the email→text conversion, e.g. `_x000D_`
# for a carriage return. Present in ~42% of descriptions.
_XML_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _decode_xml_escapes(s: str) -> str:
    def repl(m: re.Match) -> str:
        cp = int(m.group(1), 16)          # parsed as a number, never executed
        if cp in (0x0D, 0x0A):            # CR / LF → real newline
            return "\n"
        if cp == 0x09:                    # TAB → space
            return " "
        try:
            ch = chr(cp)
        except ValueError:
            return " "
        # Keep printable recovered characters; drop control chars to a space.
        return ch if ch.isprintable() else " "
    return _XML_ESCAPE.sub(repl, s)


_CID = re.compile(r"\[cid:[^\]]*\]", re.I)              # inline-image placeholder

# `visible text<mailto:…>` / `…<tel:…>` / `…<https://…>` from flattened HTML
# links — keep the visible text, drop the bracketed target (incl. giant
# exclaimer.net signature-tracking URLs).
_LINK_WRAP = re.compile(r"\s*<(?:mailto:|tel:|https?://)[^>]*>")

_BARE_URL = re.compile(r"https?://\S+")

# Conservative HTML-tag strip — only well-known tags, so we never eat a literal
# `<something>` that is real content.
_HTML_TAG = re.compile(
    r"</?(?:p|br|div|span|a|b|i|u|strong|em|ul|ol|li|table|tr|td|th|thead|tbody"
    r"|img|font|h[1-6]|blockquote|hr|pre|code)\b[^>]*>",
    re.I,
)

_MOJIBAKE = re.compile(r"\?{3,}")                       # lost-character runs → space

# ── boundary markers: content ends at the earliest of these ───────────────────
# Everything from the first match to end-of-text is signature / quoted thread /
# routing footer and is dropped.
_FOOTER = re.compile(r"\*\*\s*Created via Incoming Email Processor\s*\*\*", re.I)

# Outlook quoted-reply headers (Dutch + English). Anchored at line start so a
# mid-sentence "from:" never trips it. Verzonden/Sent/Onderwerp/Subject are
# machine-generated and never appear in normal prose.
_THREAD = re.compile(
    r"(?im)^[ \t>]*(?:van|verzonden|onderwerp|from|sent|subject)\s*:\s*\S",
)

# Sign-off phrases — the signature block always trails, so cut from here to end.
_SIGNATURE = re.compile(
    r"(?im)^[ \t]*(?:met (?:vriendelijke|hartelijke) groet(?:en)?"
    r"|vriendelijke groet(?:en)?|hartelijke groet(?:en)?"
    r"|kind regards|best regards|mvg)\b",
)

# A long rule of dashes/underscores that typically precedes a quoted block.
_SEP_RULE = re.compile(r"(?m)^[ \t]*[-_]{6,}[ \t]*$")


def _content_boundary(s: str) -> int:
    """Index where the useful content ends (start of the first sig/thread/footer
    marker), or len(s) if none is found."""
    cut = len(s)
    for pat in (_FOOTER, _THREAD, _SIGNATURE, _SEP_RULE):
        m = pat.search(s)
        if m:
            cut = min(cut, m.start())
    return cut


_WS_INLINE = re.compile(r"[ \t\f\v]+")
_WS_NEWLINES = re.compile(r"\s*\n\s*")

# Control chars openpyxl refuses to write into a worksheet cell (matches its own
# ILLEGAL_CHARACTERS_RE). Excludes \t \n \r, which are allowed. LLM output can
# carry stray control chars, so we strip these before decoding/writing.
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_NULLISH = {"nan", "none", "null", ""}


def safe_cell(value):
    """Strip worksheet-illegal control chars from a string cell (no-op otherwise).
    Used on LLM output before writing, since it bypasses `_scrub`."""
    return _ILLEGAL.sub("", value) if isinstance(value, str) else value


def _scrub(s: str, *, drop_urls: bool, cut_boundary: bool) -> str:
    """The shared scrub. `cut_boundary` truncates at the first signature/quoted-
    thread/footer marker (the aggressive, deterministic clean); omit it to keep
    that content for a downstream LLM to judge (the prepass). Newlines are
    preserved (2+ collapsed to 1) so cleaned text stays human-readable."""
    s = _ILLEGAL.sub("", s)             # drop worksheet-illegal control chars
    s = _decode_xml_escapes(s)          # `_x000D_` → newline, etc.
    if cut_boundary:
        s = s[:_content_boundary(s)]    # drop signature / quoted thread / footer
    s = _CID.sub(" ", s)                # inline-image refs
    s = _LINK_WRAP.sub("", s)           # keep link text, drop <target>
    if drop_urls:
        s = _BARE_URL.sub(" ", s)       # drop bare/tracking urls
    s = _HTML_TAG.sub(" ", s)           # known HTML tags
    s = _MOJIBAKE.sub(" ", s)           # `????` runs
    s = _WS_INLINE.sub(" ", s)          # collapse spaces/tabs
    s = _WS_NEWLINES.sub("\n", s)       # collapse blank lines, trim around newlines
    return s.strip()


def clean_text(value, *, drop_urls: bool) -> str:
    """Full deterministic clean (structural scrub + boundary cut). `drop_urls=True`
    for descriptions (urls are noise); `False` for notes (technical urls are
    signal)."""
    if value is None:
        return ""
    s = str(value)
    if s.strip().lower() in _NULLISH:
        return ""
    return _scrub(s, drop_urls=drop_urls, cut_boundary=True)


def clean_description(value) -> str:
    """Aggressive deterministic clean for `description` — the freebie/fallback
    path when the LLM cleaner (llm_clean.py) doesn't cover a ticket."""
    return clean_text(value, drop_urls=True)


def clean_note(value) -> str:
    """Light clean for time-entry `summarynotes` / `internalnotes` — keeps urls."""
    return clean_text(value, drop_urls=False)


def prepass_description(value) -> str:
    """Structural-only scrub sent to the LLM cleaner: strips encoded escapes,
    image refs, link/URL noise, HTML and mojibake to cut tokens, but KEEPS
    signatures / quoted threads / footers so the model makes that call. Returns
    "" for nullish/empty input."""
    if value is None:
        return ""
    s = str(value)
    if s.strip().lower() in _NULLISH:
        return ""
    return _scrub(s, drop_urls=True, cut_boundary=False)


# ── streaming workbook cleaner ────────────────────────────────────────────────

def clean_workbook(src: str, dst: str, targets: dict, *, sheet: str = "Export",
                   log=print) -> dict:
    """Stream `src` (.xlsx) → `dst`, applying `targets` = {column_name: fn} to the
    named columns and copying every other cell unchanged. Returns a small stats
    dict. Uses openpyxl read-only/write-only so the ~66–82MB files never fully
    materialise in memory."""
    from pathlib import Path

    from openpyxl import Workbook, load_workbook

    Path(dst).parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(src, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"{src}: no sheet named {sheet!r} (have {wb.sheetnames})")
    ws = wb[sheet]

    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))

    # Resolve target columns to (index, fn); warn (don't crash) on a missing one.
    resolved: list[tuple[int, object]] = []
    for name, fn in targets.items():
        if name in header:
            resolved.append((header.index(name), fn))
        else:
            log(f"  WARN: target column {name!r} not found in {src}")

    out = Workbook(write_only=True)
    ows = out.create_sheet(title=sheet)
    ows.append(header)

    n_rows = n_changed = 0
    for row in rows:
        row = list(row)
        for idx, fn in resolved:
            v = row[idx]
            if v is None:
                continue
            cleaned = fn(v)
            if cleaned != v:
                n_changed += 1
            row[idx] = cleaned
        ows.append(row)
        n_rows += 1
        if n_rows % 50000 == 0:
            log(f"    … {n_rows:,} rows")

    out.save(dst)
    wb.close()
    log(f"  → {dst}  ({n_rows:,} rows, {n_changed:,} cells cleaned)")
    return {"src": src, "dst": dst, "rows": n_rows, "cells_cleaned": n_changed}


def write_tickets_twin(src: str, dst: str, cleaned_map: dict, *,
                       sheet: str = "Export", log=print) -> dict:
    """Stream the raw Tickets export → cleaned twin, replacing each `description`
    with `cleaned_map[f"{instance}::{ticketnumber}"]` (the LLM output) when
    present, else the deterministic `clean_description` fallback. Key matches the
    `ticket_id` records.py builds."""
    from pathlib import Path

    from openpyxl import Workbook, load_workbook

    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"{src}: no sheet named {sheet!r} (have {wb.sheetnames})")
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    for col in (C.TicketCols.instance, C.TicketCols.ticketnumber, C.TicketCols.description):
        if col not in header:
            wb.close()
            raise SystemExit(f"{src}: missing expected column {col!r}")
    i_inst = header.index(C.TicketCols.instance)
    i_tn = header.index(C.TicketCols.ticketnumber)
    i_desc = header.index(C.TicketCols.description)

    out = Workbook(write_only=True)
    ows = out.create_sheet(title=sheet)
    ows.append(header)

    n_rows = n_llm = n_fallback = 0
    for row in rows:
        row = list(row)
        key = f"{str(row[i_inst]).strip()}::{str(row[i_tn]).strip()}"
        if key in cleaned_map:
            row[i_desc] = safe_cell(cleaned_map[key])   # LLM text bypasses _scrub
            n_llm += 1
        else:
            row[i_desc] = clean_description(row[i_desc])
            n_fallback += 1
        ows.append(row)
        n_rows += 1
        if n_rows % 50000 == 0:
            log(f"    … {n_rows:,} rows")

    out.save(dst)
    wb.close()
    log(f"  → {dst}  ({n_rows:,} rows: {n_llm:,} LLM-cleaned, {n_fallback:,} fallback)")
    return {"dst": dst, "rows": n_rows, "llm": n_llm, "fallback": n_fallback}


def clean_exports(
    tickets_xlsx: str = C.TICKETS_XLSX,
    time_entries_xlsx: str = C.TIME_ENTRIES_XLSX,
    *,
    out_tickets: str = C.CLEAN_TICKETS_XLSX,
    out_time_entries: str = C.CLEAN_TIME_ENTRIES_XLSX,
    log=print,
) -> list[dict]:
    """Clean both exports into their cleaned twins. Returns per-file stats."""
    log(f"Cleaning tickets     … {tickets_xlsx}")
    s1 = clean_workbook(
        tickets_xlsx, out_tickets,
        {C.TicketCols.description: clean_description}, log=log,
    )
    log(f"Cleaning time entries … {time_entries_xlsx}")
    s2 = clean_workbook(
        time_entries_xlsx, out_time_entries,
        {C.TimeEntryCols.summarynotes: clean_note,
         C.TimeEntryCols.internalnotes: clean_note}, log=log,
    )
    return [s1, s2]


if __name__ == "__main__":
    import argparse

    ap = argparse.ArgumentParser(description="Clean ticket text (Section 1.1).")
    ap.add_argument("--tickets", default=C.TICKETS_XLSX)
    ap.add_argument("--time-entries", default=C.TIME_ENTRIES_XLSX)
    ap.add_argument("--out-tickets", default=C.CLEAN_TICKETS_XLSX)
    ap.add_argument("--out-time-entries", default=C.CLEAN_TIME_ENTRIES_XLSX)
    args = ap.parse_args()

    stats = clean_exports(
        args.tickets, args.time_entries,
        out_tickets=args.out_tickets, out_time_entries=args.out_time_entries,
    )
    print(f"Done: {stats}")
