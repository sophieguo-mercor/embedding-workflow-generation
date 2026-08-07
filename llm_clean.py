#!/usr/bin/env python3
"""
Section 1.1 (LLM variant) — clean ticket DESCRIPTIONS with the Anthropic Message
Batches API, then write the cleaned Tickets twin. Notes stay deterministic
(clean.py); only the description column — a noisy HTML-to-text email dump — is
handed to the model.

Why an LLM here (vs. the regex in clean.py): the ~8% of descriptions with genuine
quoted reply threads and the endless variety of signature layouts are exactly the
fuzzy cases regex boundaries get wrong. The model keeps the requester's own
wording (no translation, no summarising) and drops the noise.

Why the Batch API (mirrors resolution-extraction/extract.py): the corpus is
~140k descriptions → ~11–12k requests at 12/req. The batch queue is a separate,
higher-throughput lane at 50% the token price — you trade latency (minutes to a
few hours, no SLA) for throughput and cost.

Cost control:
  * A deterministic prepass (clean.prepass_description) strips encoded escapes,
    image refs, and giant tracking URLs BEFORE sending — fewer tokens, but the
    signature/thread judgment is left to the model.
  * Empty / trivially-short descriptions are never sent (freebie path); the twin
    writer falls back to the deterministic clean for them.
  * The static system prompt (instructions + few-shot) is cache_control'd, billed
    once and cache-read (~0.1x) for the rest of the batch.

Resume-safe on two levels (like extract.py):
  1. cache/desc_clean.jsonl — ticket ids already cleaned are skipped on re-runs.
  2. cache/desc_batch_manifest.json — an in-flight batch is reconnected, never
     re-submitted (double-spend guard).

CLI
---
    # Free: print the prompt + a request/token estimate, no API calls.
    python llm_clean.py --dry-run

    # Small real batch (a few cents) to eyeball output quality.
    python llm_clean.py --sample 24

    # Full run. Submit and walk away; re-run (or --collect) to reconnect + finish.
    python llm_clean.py
    python llm_clean.py --submit-only      # create batch + exit
    python llm_clean.py --collect          # collect an already-submitted batch
"""
from __future__ import annotations

import argparse
import json
import os
import re
import time
from pathlib import Path

import clean
import config as C

MAX_REQUESTS_PER_BATCH = 100_000     # Anthropic hard limit (we're far under)
MIN_CONTENT_CHARS = 40               # shorter post-prepass → not worth an API call


def load_dotenv(path: str = ".env") -> None:
    """Minimal .env loader — no dependency. Existing env vars win."""
    p = Path(path)
    if not p.exists():
        return
    for line in p.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, val = line.partition("=")
        os.environ.setdefault(key.strip(), val.strip().strip("'\""))


# ── prompt ────────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """You clean IT-support ticket DESCRIPTIONS for a Dutch MSP. Each description is the customer's original request as it arrived — usually by email — so the real content is buried in noise: signatures, quoted earlier email threads, automated footers, greetings, sign-offs, image placeholders, and tracking links.

Return ONLY the substantive content of the request: what the customer wants, or what is broken.

KEEP:
- The actual request or problem statement (the ask, the symptom).
- The affected system / application / license, device, mailbox, server, or file path.
- The person or account the request is about (e.g. "toegang voor Lisanne").
- Concrete details: error messages, product names, hostnames, addresses that are part of the request.

REMOVE:
- Signature blocks (name, job title, company, postal address, phone, website, social links).
- Quoted earlier email threads: everything from a header line like "Van:" / "Verzonden:" / "Onderwerp:" or "From:" / "Sent:" / "Subject:" onward, and any reply history.
- Automated footers ("**Created via Incoming Email Processor**", "From:"/"To:" routing lines).
- Greetings and sign-offs ("Goedemiddag,", "Met vriendelijke groet,", "Bedankt!").
- Image placeholders, tracking URLs, and leftover markup.

RULES:
- Keep the original language (Dutch or English) — do NOT translate.
- Do NOT summarise, rephrase, or add words. Return the customer's own wording with the noise removed; you may lightly join lines broken mid-sentence.
- If nothing substantive remains after removing noise, return an empty string "".

## Output
Return ONLY a JSON array — one object per input ticket, same order, no markdown fences, no prose:
[{"id": 0, "text": "..."}, {"id": 1, "text": ""}]

## Examples
Input:
--- TICKET 0 ---
Goedemiddag, Graag zou ik voor Lisanne van 't Sant toegang willen aanvragen tot de mailbox aanbestedingen@gzicht.nl. Kunnen jullie dit verwerken? Vriendelijke groet, Lindsy Fredriksz office manager 06 48 52 37 44 Stationsweg 73C, 6711 PL Ede
--- TICKET 1 ---
Beste, Bedankt voor uw mail. Met vriendelijke groet, Jan de Vries Directeur | ACME BV www.acme.nl
--- TICKET 2 ---
Ik kan sinds de laatste update geen verbinding meer maken met de netwerkschijven, de VPN is wel verbonden. Met vriendelijke groet, Christian Kalisvaart Adviseur Van: support Verzonden: dinsdag Onderwerp: RE: netwerk
Output:
[{"id": 0, "text": "Graag zou ik voor Lisanne van 't Sant toegang willen aanvragen tot de mailbox aanbestedingen@gzicht.nl. Kunnen jullie dit verwerken?"}, {"id": 1, "text": ""}, {"id": 2, "text": "Ik kan sinds de laatste update geen verbinding meer maken met de netwerkschijven, de VPN is wel verbonden."}]
"""


def build_user_message(descriptions: list[str], char_limit: int) -> str:
    parts = []
    for i, d in enumerate(descriptions):
        text = (d or "").strip()
        if len(text) > char_limit:
            text = text[:char_limit] + " …[truncated]"
        parts.append(f"--- TICKET {i} ---\n{text}")
    return "\n\n".join(parts)


def _strip_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def parse_result_text(text: str, n: int) -> dict[int, str]:
    """Parse one succeeded response → {position: cleaned_text}. Raises on bad JSON
    so the caller can skip the group and re-submit it on a later run."""
    parsed = json.loads(_strip_fences(text))
    out: dict[int, str] = {}
    for item in parsed:
        pos = item.get("id")
        if isinstance(pos, int) and 0 <= pos < n:
            out[pos] = str(item.get("text", "")).strip()
    return out


# ── batch client ──────────────────────────────────────────────────────────────

class DescriptionCleaner:
    """Builds batch Requests. No network in the constructor beyond client init, so
    the same object serves dry-run, submit, and collect paths."""

    def __init__(self, model: str = C.BATCH_MODEL, max_tokens: int = C.BATCH_MAX_TOKENS):
        self.model = model
        self.max_tokens = max_tokens
        try:
            import anthropic  # noqa: F401
        except ImportError:
            raise SystemExit("pip install anthropic")
        from anthropic import Anthropic
        from anthropic.types.message_create_params import MessageCreateParamsNonStreaming
        from anthropic.types.messages.batch_create_params import Request

        self._MCP = MessageCreateParamsNonStreaming
        self._Request = Request
        if not os.environ.get("ANTHROPIC_API_KEY"):
            raise SystemExit("Set ANTHROPIC_API_KEY in your environment (or .env).")
        self.client = Anthropic()

    def make_request(self, custom_id: str, descriptions: list[str], char_limit: int):
        return self._Request(
            custom_id=custom_id,
            params=self._MCP(
                model=self.model,
                max_tokens=self.max_tokens,
                system=[{"type": "text", "text": SYSTEM_PROMPT,
                         "cache_control": {"type": "ephemeral"}}],
                messages=[{"role": "user",
                           "content": build_user_message(descriptions, char_limit)}],
            ),
        )


# ── ticket reading ────────────────────────────────────────────────────────────

def read_descriptions(tickets_xlsx: str, *, sheet: str = "Export", log=print):
    """Stream the raw Tickets export → [(ticket_id, prepassed_description)], one per
    unique (instance, ticketnumber). Empty/short prepass results are dropped here
    (freebie path) and never sent to the API."""
    from openpyxl import load_workbook

    wb = load_workbook(tickets_xlsx, read_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    i_inst = header.index(C.TicketCols.instance)
    i_tn = header.index(C.TicketCols.ticketnumber)
    i_desc = header.index(C.TicketCols.description)

    items: list[tuple[str, str]] = []
    seen: set[str] = set()
    n_rows = n_empty = 0
    for row in rows:
        n_rows += 1
        key = f"{str(row[i_inst]).strip()}::{str(row[i_tn]).strip()}"
        if key in seen:
            continue
        seen.add(key)
        pp = clean.prepass_description(row[i_desc])
        if len(pp) < MIN_CONTENT_CHARS:
            n_empty += 1
            continue
        items.append((key, pp))
    wb.close()
    log(f"  {n_rows:,} rows | {len(seen):,} unique tickets | "
        f"{len(items):,} to clean | {n_empty:,} empty/short (freebie)")
    return items


# ── cache + manifest ──────────────────────────────────────────────────────────

def load_cache(path: str) -> dict[str, str]:
    out: dict[str, str] = {}
    p = Path(path)
    if p.exists():
        with p.open(encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    rec = json.loads(line)
                    out[rec["id"]] = rec.get("text", "")
                except Exception:
                    continue
    return out


class CacheWriter:
    """Append-mode, idempotent {id,text} sink for cleaned descriptions."""

    def __init__(self, path: str, done: set[str]):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self._fh = open(path, "a", encoding="utf-8")
        self._done = done

    def emit(self, ticket_id: str, text: str) -> bool:
        if ticket_id in self._done:
            return False
        self._fh.write(json.dumps({"id": ticket_id, "text": text}, ensure_ascii=False) + "\n")
        self._fh.flush()
        self._done.add(ticket_id)
        return True

    def close(self):
        self._fh.close()


def build_groups(items, done: set[str], batch_size: int) -> dict[str, dict]:
    """Chunk not-yet-cleaned items into {custom_id: {ids, descriptions}}."""
    todo = [(k, d) for k, d in items if k not in done]
    groups: dict[str, dict] = {}
    for c, i in enumerate(range(0, len(todo), batch_size)):
        chunk = todo[i:i + batch_size]
        groups[f"g{c:06d}"] = {"ids": [k for k, _ in chunk],
                               "descriptions": [d for _, d in chunk]}
    return groups


def save_manifest(path: str, batch_id: str, model: str, groups: dict) -> None:
    slim = {cid: {"ids": g["ids"]} for cid, g in groups.items()}   # drop text
    Path(path).write_text(
        json.dumps({"batch_id": batch_id, "model": model, "groups": slim},
                   ensure_ascii=False), encoding="utf-8")


# ── submit / collect ──────────────────────────────────────────────────────────

def collect(cleaner: DescriptionCleaner, manifest_path: str, cache_path: str,
            poll_interval: int, *, log=print) -> None:
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    batch_id = manifest["batch_id"]
    groups = manifest["groups"]
    log(f"Reconnecting to batch {batch_id} ({len(groups):,} requests) …")

    while True:
        b = cleaner.client.messages.batches.retrieve(batch_id)
        c = b.request_counts
        log(f"  status={b.processing_status}  processing={c.processing} "
            f"succeeded={c.succeeded} errored={c.errored} "
            f"canceled={c.canceled} expired={c.expired}")
        if b.processing_status == "ended":
            break
        time.sleep(poll_interval)

    done = set(load_cache(cache_path))
    writer = CacheWriter(cache_path, done)
    written = errored = bad_json = 0
    usage = {"input": 0, "output": 0, "cache_read": 0, "cache_create": 0}

    for r in cleaner.client.messages.batches.results(batch_id):
        grp = groups.get(r.custom_id)
        if grp is None:
            continue
        ids = grp["ids"]
        if r.result.type != "succeeded":
            err = getattr(getattr(r.result, "error", None), "type", r.result.type)
            errored += len(ids)
            log(f"    ! {r.custom_id} {err} ({len(ids)}) — re-submit next run")
            continue
        msg = r.result.message
        u = msg.usage
        usage["input"] += getattr(u, "input_tokens", 0) or 0
        usage["output"] += getattr(u, "output_tokens", 0) or 0
        usage["cache_read"] += getattr(u, "cache_read_input_tokens", 0) or 0
        usage["cache_create"] += getattr(u, "cache_creation_input_tokens", 0) or 0
        text = "".join(blk.text for blk in msg.content if blk.type == "text")
        try:
            by_pos = parse_result_text(text, len(ids))
        except Exception as e:
            bad_json += len(ids)
            log(f"    ! {r.custom_id} bad JSON ({len(ids)}): {e} — re-submit next run")
            continue
        for pos, tid in enumerate(ids):
            if writer.emit(tid, by_pos.get(pos, "")):
                written += 1
    writer.close()
    Path(manifest_path).unlink(missing_ok=True)

    log(f"\nCollected {batch_id}: {written:,} cleaned, {errored:,} errored, "
        f"{bad_json:,} unparseable.")
    log(f"Tokens — input {usage['input']:,} (cache read {usage['cache_read']:,}, "
        f"create {usage['cache_create']:,}) / output {usage['output']:,}")
    if errored or bad_json:
        log("Some tickets weren't cleaned; re-run to submit a fresh batch for them.")


def submit(cleaner: DescriptionCleaner, groups: dict, manifest_path: str, *,
           char_limit: int, log=print) -> None:
    if len(groups) > MAX_REQUESTS_PER_BATCH:
        raise SystemExit(f"{len(groups):,} requests exceeds the "
                         f"{MAX_REQUESTS_PER_BATCH:,}/batch limit. Use --sample.")
    requests = [cleaner.make_request(cid, g["descriptions"], char_limit)
                for cid, g in groups.items()]
    log(f"Submitting batch of {len(requests):,} requests …")
    batch = cleaner.client.messages.batches.create(requests=requests)
    save_manifest(manifest_path, batch.id, cleaner.model, groups)
    log(f"  batch {batch.id}  status={batch.processing_status}  (manifest → {manifest_path})")


# ── stage entry point ─────────────────────────────────────────────────────────

def run_clean_stage(args, *, log=print) -> None:
    """Full cleaning stage: notes twin (deterministic) + description twin (LLM)."""
    load_dotenv()
    tickets_xlsx = getattr(args, "tickets", C.TICKETS_XLSX)
    time_entries_xlsx = getattr(args, "time_entries", C.TIME_ENTRIES_XLSX)
    manifest = C.DESC_BATCH_MANIFEST
    cache = C.DESC_CLEAN_CACHE
    batch_size = getattr(args, "batch_size", None) or C.BATCH_SIZE
    char_limit = getattr(args, "char_limit", None) or C.BATCH_CHAR_LIMIT
    sample = getattr(args, "sample", None) or getattr(args, "limit", None)
    poll = getattr(args, "poll_interval", None) or 30

    # ── collect-only / reconnect an in-flight batch (double-spend guard) ──────
    if getattr(args, "collect", False):
        if not Path(manifest).exists():
            raise SystemExit(f"No {manifest} to collect. Submit a batch first.")
        collect(DescriptionCleaner(), manifest, cache, poll, log=log)
    elif Path(manifest).exists() and not getattr(args, "force_new", False):
        log(f"Found in-flight batch in {manifest} — resuming it "
            f"(use --force-new to discard).")
        collect(DescriptionCleaner(), manifest, cache, poll, log=log)
    else:
        # ── build requests locally ────────────────────────────────────────────
        log(f"Reading descriptions … {tickets_xlsx}")
        items = read_descriptions(tickets_xlsx, log=log)
        if sample:
            items = items[:sample]
            log(f"  --sample/--limit → first {len(items):,} tickets")
        done = set(load_cache(cache))
        groups = build_groups(items, done, batch_size)
        log(f"  {len(groups):,} requests to submit "
            f"({len(done):,} already cleaned, skipped)")

        if getattr(args, "dry_run", False):
            log("\n===== SYSTEM PROMPT (cached prefix) =====\n" + SYSTEM_PROMPT[:1600])
            if groups:
                first = next(iter(groups.values()))
                log("\n===== USER MESSAGE (first request) =====\n"
                    + build_user_message(first["descriptions"], char_limit)[:2000])
            approx_in = sum(len(d) for _, d in items) // 4
            log(f"\n===== PLAN =====\nrequests: {len(groups):,}  "
                f"~input tokens (uncached user): {approx_in:,}  "
                f"(model {C.BATCH_MODEL}, batch_size {batch_size}) — NO API CALLS")
            return

        if getattr(args, "force_new", False):
            Path(manifest).unlink(missing_ok=True)
        if not groups:
            log("Nothing to submit — all descriptions already cleaned or empty.")
        else:
            cleaner = DescriptionCleaner()
            submit(cleaner, groups, manifest, char_limit=char_limit, log=log)
            if getattr(args, "submit_only", False):
                log("Submitted. Collect later with:  python llm_clean.py --collect")
                return
            collect(cleaner, manifest, cache, poll, log=log)

    # ── write both cleaned twins ──────────────────────────────────────────────
    if getattr(args, "dry_run", False):
        return
    cache_map = load_cache(cache)
    log(f"\nWriting cleaned twins ({len(cache_map):,} LLM-cleaned descriptions)…")
    clean.clean_workbook(
        time_entries_xlsx, C.CLEAN_TIME_ENTRIES_XLSX,
        {C.TimeEntryCols.summarynotes: clean.clean_note,
         C.TimeEntryCols.internalnotes: clean.clean_note}, log=log)
    clean.write_tickets_twin(tickets_xlsx, C.CLEAN_TICKETS_XLSX, cache_map, log=log)
    log(f"→ cleaned twins in {C.CLEANED_DIR}/  (consume with --use-cleaned)")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tickets", default=C.TICKETS_XLSX)
    ap.add_argument("--time-entries", dest="time_entries", default=C.TIME_ENTRIES_XLSX)
    ap.add_argument("--sample", type=int, default=None, help="clean only the first N tickets")
    ap.add_argument("--batch-size", type=int, default=C.BATCH_SIZE)
    ap.add_argument("--char-limit", type=int, default=C.BATCH_CHAR_LIMIT)
    ap.add_argument("--poll-interval", type=int, default=30)
    ap.add_argument("--submit-only", action="store_true", help="create the batch and exit")
    ap.add_argument("--collect", action="store_true", help="collect an already-submitted batch")
    ap.add_argument("--force-new", action="store_true", help="discard an in-flight manifest")
    ap.add_argument("--dry-run", action="store_true", help="print prompt + estimate, no API calls")
    args = ap.parse_args()
    run_clean_stage(args)


if __name__ == "__main__":
    main()
